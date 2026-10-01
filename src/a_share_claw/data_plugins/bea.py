"""Keyless BEA PCE release facts; reported percentages, never reconstructed indices."""
from __future__ import annotations

import calendar
import io
import re
from datetime import datetime, datetime as NativeDateTime
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

from .core import DataError, Manifest, Payload, Provider, iso_date
from .normalization import _selection, checked_result, numeric
from .providers import _Publication, parameters

METRICS=("pce_price_mom_reported", "core_pce_price_mom_reported", "pce_price_yoy_reported", "core_pce_price_yoy_reported")
MONTHS=tuple(calendar.month_name)[1:]
NOTES=["BEA-reported percent changes, not raw price-index levels or core recomputed rates",
       "Current release capture only; original release clock and capture clock are distinct; historical revision vintage is not certified"]


def native_requirement(metric,year,month,fact_id):
    if metric not in METRICS or type(year) is not int or not 1<=year<=9999 or type(month) is not int or not 1<=month<=12:
        raise DataError("invalid_request","Pin a supported BEA reported PCE rate and explicit year/month")
    start=f"{year:04d}-{month:02d}-01";end=f"{year:04d}-{month:02d}-{calendar.monthrange(year,month)[1]:02d}"
    return {"fact_id":fact_id,"entity":"US","metric":metric,"unit":"percent","value_type":"number",
            "data_period":start+"/"+end,"observation_start":end,"observation_end":end}


class Release(_Publication):
    def __init__(self):
        super().__init__();self.titles=[];self._title=None
    def handle_starttag(self,tag,attrs):
        super().handle_starttag(tag,attrs)
        if tag=="h1":self._title=[]
    def handle_data(self,data):
        super().handle_data(data)
        if self._title is not None:self._title.append(data)
    def handle_endtag(self,tag):
        super().handle_endtag(tag)
        if tag=="h1" and self._title is not None:
            self.titles.append(" ".join("".join(self._title).split()));self._title=None


def release_values(text,month):
    """Require paired headline/core native statements within the same named month."""
    values={};excerpts={}
    number=r"(?P<value>\d+(?:\.\d+)?)"
    for label,suffix,tail in (("preceding month","mom",r""),("same month one year ago","yoy",r" from one year ago")):
        pattern=(rf"From the {label}, the PCE price index for {re.escape(MONTHS[month-1])} "
                 rf"(?P<head_direction>increased|decreased) (?P<head>\d+(?:\.\d+)?) percent\. "
                 rf"Excluding food and energy, the PCE price index (?P<core_direction>increased|decreased) {number} percent{tail}\.")
        matches=list(re.finditer(pattern,text))
        if not matches:raise DataError("insufficient_coverage","Exact monthly headline/core PCE rate statements are absent")
        pairs={(float(m["head"])*(-1 if m["head_direction"]=="decreased" else 1),
                float(m["value"])*(-1 if m["core_direction"]=="decreased" else 1)) for m in matches}
        if len(pairs)!=1:raise DataError("source_disagreement","Repeated native PCE statements disagree")
        head,core=next(iter(pairs))
        for key,value in ((f"pce_price_{suffix}_reported",head),(f"core_pce_price_{suffix}_reported",core)):
            values[key]=numeric(value);excerpts[key]=list(dict.fromkeys(m.group() for m in matches))
    return values,excerpts


class BEA(Provider):
    manifest=Manifest("bea",("macro.pce_release_snapshot","macro.pce_history_snapshot"),("www.bea.gov",),version="1.1.0")
    def fetch(self,r):
        if r.capability=="macro.pce_history_snapshot":return self.fetch_history(r)
        p=parameters(r,{"url","year","month"},{"url","year","month"})
        native=native_requirement(METRICS[0],p["year"],p["month"],"validation")
        if r.as_of_date!=datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat():
            raise DataError("historical_unavailable","BEA release snapshots require today's capture date")
        url=p["url"];u=urlsplit(url)
        if (u.scheme!="https" or u.hostname!="www.bea.gov" or u.username or u.password or u.port not in (None,443) or u.query or u.fragment or
                not re.fullmatch(r"/news/\d{4}/personal-income-and-outlays-"+MONTHS[p["month"]-1].lower()+rf"-{p['year']}",u.path)):
            raise DataError("source_denied","Use the exact official BEA release URL for the planned month/year")
        raw=self.transport.get(url);page=raw.decode("utf-8-sig")
        if "<html" not in page.lower():raise DataError("invalid_schema","BEA release must be HTML")
        parser=Release();parser.feed(page)
        title=f"Personal Income and Outlays, {MONTHS[p['month']-1]} {p['year']}"
        if [t for t in parser.titles if t!="News Release"]!=[title]:raise DataError("period_mismatch","Official release heading must match the exact planned month/year")
        text=" ".join("".join(parser.text).split())
        clocks=list(re.finditer(r"EMBARGOED UNTIL RELEASE AT (\d{1,2}):(\d{2}) (a\.m\.|p\.m\.) (EDT|EST), ([A-Za-z]+), ([A-Za-z]+) (\d{1,2}), (\d{4})",text))
        if len(clocks)!=1:raise DataError("missing_publication_date","Require one explicitly labelled BEA release clock")
        m=clocks[0];hour=int(m[1]);minute=int(m[2])
        if not 1<=hour<=12:raise DataError("invalid_schema","Invalid publisher clock")
        released=datetime(int(m[8]),MONTHS.index(m[6])+1,int(m[7]),hour%12+(12 if m[3]=="p.m." else 0),minute,tzinfo=ZoneInfo("America/New_York"))
        if released.tzname()!=m[4] or released.strftime("%A")!=m[5]:raise DataError("invalid_schema","Publisher date/weekday/timezone disagree")
        published=released.date().isoformat()
        if u.path.split("/")[2]!=str(released.year):raise DataError("source_mismatch","Release URL year and publisher clock disagree")
        if published>r.as_of_date or native["observation_end"]>published or released>datetime.now(released.tzinfo):
            raise DataError("future_data","Release/observation period exceeds capture cutoff")
        values,excerpts=release_values(text,p["month"])
        return Payload({"title":title,"year":p["year"],"month":p["month"],"values":values,"source_excerpts":excerpts,
            "unit":"percent","publication_date":published,"publisher_available_at":released.isoformat(),"publication_time_precision":"minute",
            "snapshot_as_of_date":r.as_of_date,"availability_basis":"observed_current_snapshot","historical_vintage_certified":False,
            "source_notes":NOTES},raw,url,"verified",NOTES)

    def fetch_history(self,r):
        """Native dated comparisons from one published workbook; no interpolation."""
        from openpyxl import load_workbook
        p=parameters(r,{"url","year","month"},{"url","year","month"})
        native_requirement(METRICS[0],p["year"],p["month"],"validation")
        if r.as_of_date!=datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat():
            raise DataError("historical_unavailable","Historical observations still require today's capture date")
        u=urlsplit(p["url"])
        pattern=rf"/sites/default/files/\d{{4}}-\d{{2}}/pi{p['month']:02d}{p['year']%100:02d}-hist\.xlsx"
        if (u.scheme!="https" or u.hostname!="www.bea.gov" or u.username or u.password or u.port not in (None,443) or u.query or u.fragment or not re.fullmatch(pattern,u.path)):
            raise DataError("source_denied","Use the exact official BEA historical-comparison workbook")
        raw=self.transport.get(p["url"])
        try:
            w=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
            if w.sheetnames!=["PIOhist_M"]:raise ValueError()
            rows=list(w.active.values);w.close()
            published=rows[0][6]
            if (not isinstance(published,NativeDateTime) or rows[1][0]!=f"{MONTHS[p['month']-1]} {p['year']} Personal Income and Outlays" or
                    rows[2][0]!="Historical Comparisons" or rows[4][1]!=datetime(p["year"],p["month"],1) or
                    tuple(rows[4][2::2])!=("Last period with equal value","Last period with larger value","Last period with smaller value") or
                    u.path.split("/")[4]!=published.strftime("%Y-%m")):
                raise ValueError()
            publication=published.date().isoformat()
            if publication>r.as_of_date or native_requirement(METRICS[0],p['year'],p['month'],'v')['observation_end']>publication:
                raise DataError("future_data","Workbook publication/observation exceeds capture")
            values={};excerpts={};section=None;price=False
            for row in rows:
                if row[0]=="Chain-type price indexes":price=True;section=None;continue
                if not price:continue
                if row[0] in {"Percent change from preceding month:","Percent change from month one year ago:"}:
                    section="mom" if row[0]=="Percent change from preceding month:" else "yoy";continue
                if row[0] not in {"PCE","PCE, excluding food and energy"}:section=None;continue
                if section is None:continue
                metric=("core_" if row[0].startswith("PCE,") else "")+"pce_price_"+section+"_reported"
                pairs=[(f"{p['year']}M{p['month']:02d}",row[1])]+[(row[i],row[i+1]) for i in (2,4,6) if row[i]!="---"]
                for period,value in pairs:
                    if not isinstance(period,str) or not re.fullmatch(r"\d{4}M(0[1-9]|1[0-2])",period):raise ValueError()
                    y,m=int(period[:4]),int(period[5:]);n=native_requirement(metric,y,m,'v')
                    if n['observation_end']>publication:raise DataError("future_data","Future comparison observation")
                    key=metric+"."+period
                    v=numeric(value)
                    if period!=f"{p['year']}M{p['month']:02d}":
                        current=numeric(row[1])
                        for i,relation in ((2,lambda a,b:a==b),(4,lambda a,b:a>b),(6,lambda a,b:a<b)):
                            if row[i]==period and not relation(numeric(row[i+1]),current):
                                raise DataError("source_disagreement","Workbook equal/larger/smaller labels contradict values")
                    if key in values and values[key]!=v:raise DataError("source_disagreement","Workbook dated comparisons disagree")
                    values[key]=v;excerpts[key]=[str(row),"Printed dated value; not a continuous series"]
            if any(metric+f".{p['year']}M{p['month']:02d}" not in values for metric in METRICS):raise ValueError()
        except DataError:raise
        except (ValueError,TypeError,KeyError,IndexError):raise DataError("invalid_schema","Unsupported BEA comparison workbook") from None
        notes=NOTES+["All selected comparison values share this workbook publication version; sparse last-equal/larger/smaller comparisons are not a continuous history"]
        return Payload({"year":p['year'],"month":p['month'],"values":values,"source_excerpts":excerpts,"unit":"percent",
            "publication_date":publication,"publisher_available_at":None,"publication_time_precision":"day",
            "snapshot_as_of_date":r.as_of_date,"availability_basis":"observed_current_snapshot","historical_vintage_certified":False,"source_notes":notes},raw,p['url'],"verified",notes)


def reported_observation(result,*,metric,year,month,period_kind):
    capability=result["capability"]
    if capability not in {"macro.pce_release_snapshot","macro.pce_history_snapshot"}:raise DataError("source_mismatch","Unsupported BEA capability")
    data=checked_result(result,"bea",capability)
    history=capability=="macro.pce_history_snapshot"
    key=metric+f".{year}M{month:02d}" if history else metric
    native=native_requirement(metric,year,month,"selection")
    if period_kind!="month" or (not history and (data.get("year")!=year or data.get("month")!=month)) or data.get("unit")!="percent" or key not in data["values"]:
        raise DataError("period_mismatch","BEA reported rates retain exact monthly identity and native percent units")
    cutoff=result["provenance"]["as_of_date"]
    capture=datetime.fromisoformat(result["provenance"]["retrieved_at"])
    release=datetime.fromisoformat(data["publisher_available_at"]) if data["publisher_available_at"] is not None else None
    if (capture.tzinfo is None or (release is not None and release.tzinfo is None) or capture.astimezone(ZoneInfo("Asia/Shanghai")).date().isoformat()!=cutoff or
            data.get("snapshot_as_of_date")!=cutoff or data.get("availability_basis")!="observed_current_snapshot" or
            data.get("historical_vintage_certified") is not False or (release is not None and (release>capture or release.date().isoformat()!=iso_date(data["publication_date"]))) or data["publication_date"]>cutoff or native["observation_end"]>data["publication_date"]):
        raise DataError("future_data","BEA release and actual capture clocks must remain eligible and distinct")
    return _selection("bea",{"metric":metric,"value":numeric(data["values"][key]),"unit":"percent","period_kind":"month",
        "period_start":native["data_period"].split("/")[0],"period_end":native["observation_end"],
        "publication_date":data["publication_date"],"publisher_available_at":data["publisher_available_at"],
        "publication_time_precision":data["publication_time_precision"],"available_at":capture.isoformat(),"snapshot_as_of_date":cutoff,
        "eligibility":"verified_current_snapshot","availability_basis":"observed_current_snapshot","historical_vintage_certified":False,
        "mapping_version":"bea-reported-pce-v1","source_excerpts":data["source_excerpts"][key],"source_notes":data["source_notes"]},[result])
