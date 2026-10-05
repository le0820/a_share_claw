"""Explicit native source bindings for the frozen five-chart research view."""
from pathlib import Path
import json,re,calendar,hashlib,io,math
from datetime import datetime,timezone
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from .providers import _Publication
from ..harness.ai_pressure import compute
from ..harness.five_charts import CHARTS

SW={'食品饮料':'801120','家用电器':'801110','商贸零售':'801200','纺织服饰':'801130','轻工制造':'801140','社会服务':'801210','美容护理':'801980','汽车':'801880','通信':'801770','半导体':'801081','银行':'801780'}

def text(raw):
 p=_Publication();p.feed(raw.decode('utf-8-sig',errors='replace'));return p,' '.join(''.join(p.text).split())

def dc_rows(sources,read):
 rows=[]
 for s in sources:
  m=re.fullmatch(r'(?:nvda|news)_dc_(\d{4})_q([1-4])',s['id'])
  if not m:continue
  _,t=text(read(s['id']));value=re.search(r'(?:Record (?:quarterly )?)?Data Center revenue (?:of |was (?:a record )?)?\$([\d.]+) billion',t,re.I)
  if value is None:value=re.search(r'Data Center (?:First|Second|Third|Fourth)-quarter revenue was (?:a record )?\$([\d.]+) billion',t,re.I)
  pm=re.search(r'(January|February|March|April|May|June|July|August|September|October|November|December) (\d{1,2}), (\d{4})',t)
  end=re.search(r'(?:quarter ended|ended) ([A-Za-z]+) (\d{1,2}), (\d{4})',t)
  if not(value and pm and end):continue
  def day(m):return f'{int(m[3]):04}-{list(calendar.month_name).index(m[1]):02}-{int(m[2]):02}'
  rows.append({'fiscal_year':int(m[1]),'quarter':int(m[2]),'release_date':day(pm),'period_end':day(end),'value_usd_millions':float(value[1])*1000,'precision':'native rounded issuer headline','source_id':s['id']})
 unique={}
 for r in rows:
  key=(r['fiscal_year'],r['quarter'])
  if key in unique and unique[key]['value_usd_millions']!=r['value_usd_millions']:raise ValueError('issuer_dc_source_disagreement')
  unique[key]=r
 return sorted(unique.values(),key=lambda r:(r['fiscal_year'],r['quarter']))

def normalize(sources,read,*,scope_key,run_id,as_of_date,generated_at=None):
 names={s['id'] for s in sources};series=[];notes=[];errors={}
 def add(name,points,unit,basis,**kw):
  points=sorted(points,key=lambda r:r['date']);unique={}
  for r in points:
   if r['date'] in unique and r['value']!=unique[r['date']]['value']:raise ValueError('duplicate_observation_disagreement:'+name)
   unique[r['date']]=r
  series.append({'name':name,'points':list(unique.values()),'unit':unit,'basis':basis,**kw})
 def price(name,rows,source,identity,basis):
  rows=sorted(rows,key=lambda r:r[0]);rows=[(d,float(v)) for d,v in rows if '2025-01-01'<=d<=as_of_date]
  if not rows or len(set(d for d,v in rows))!=len(rows) or any(not math.isfinite(v) or v<=0 for d,v in rows):raise ValueError('invalid_native_price:'+name)
  if rows[0][0]!='2025-01-02':raise ValueError('price_window_missing_first_session')
  base=rows[0][1];add(name,[{'date':d,'value':v/base*100,'native_value':v,'source_ids':[source]} for d,v in rows],'price_first_observation_100',basis+'；'+identity+'；各自首个观测='+rows[0][0]+'，不作汇率或分红调整')
 catalog={}
 if 'sw_catalog_valid' in names:
  c=json.loads(read('sw_catalog_valid'))['data']
  if c['count']!=31 or len(c['results'])!=31:raise ValueError('incomplete_native_level_one_catalog')
  catalog={r['swindexcode']:r['swindexname'] for r in c['results']}
 secondary={}
 if 'sw_secondary_catalog_retry' in names:
  c=json.loads(read('sw_secondary_catalog_retry'))['data']
  if c['count']!=len(c['results']):raise ValueError('incomplete_native_secondary_catalog')
  secondary={r['swindexcode']:r['swindexname'] for r in c['results']}
 for name,code in SW.items():
  sid='sw_semiconductor_retry' if name=='半导体' and 'sw_'+code not in names else 'sw_'+code
  try:
   identity_map=secondary if name=='半导体' else catalog
   if identity_map.get(code)!=name:raise ValueError('native_sw_catalog_identity_required')
   rows=json.loads(read(sid))['data']
   if any(str(r['swindexcode'])!=code for r in rows):raise ValueError('wrong_native_sw_identity')
   price(name,[(r['bargaindate'],r['closeindex']) for r in rows],sid,code,('申万二级行业' if name=='半导体' else '申万一级行业'))
  except (KeyError,ValueError,TypeError) as e:errors[name]=str(e)
 for name,code,market,native in [('恒生科技','HZ5017',27,'恒生科技指数'),('费城半导体','A_SOX',12,'纳指费城半导体'),('纳斯达克100','A_NDX',12,'纳斯达克100'),('标普500','A_SPX',12,'标普500')]:
  sid='tdx_'+code
  try:
   d=json.loads(read(sid));matches=[r for r in d['identity'] if r['market']==market and r['code']==code and r['name']==native]
   if len(matches)!=1:raise ValueError('native_identity_not_unique')
   price(name,[(r['datetime'][:10],r['close']) for r in d['bars']],sid,f'easy-tdx market={market}, code={code}',native)
  except (KeyError,ValueError,TypeError) as e:errors[name]=str(e)
 ns={'d':'http://schemas.microsoft.com/ado/2007/08/dataservices','m':'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'};real={};nominal={};cn={}
 for sid in names:
  if sid.startswith('treasury_real_') or sid.startswith('treasury_nominal_'):
   target=real if sid.startswith('treasury_real_') else nominal;field='TC_10YEAR' if target is real else 'BC_10YEAR'
   for p in ET.fromstring(read(sid)).findall('.//m:properties',ns):
    v=p.find('d:'+field,ns);d=p.find('d:NEW_DATE',ns).text[:10]
    if v is not None and v.text is not None and d<=as_of_date:target[d]={'date':d,'value':float(v.text),'source_ids':[sid]}
  if sid.startswith('china_curve_') and sid!='china_curve_page':
   for r in json.loads(read(sid))['heList']:
    d=r['workTime'][:10]
    if d<=as_of_date and r.get('tenYear') not in (None,'','--'):cn[d]={'date':d,'value':float(r['tenYear']),'source_ids':[sid]}
 if real:add('10年TIPS',[r for d,r in sorted(real.items()) if d>='2025-01-01'],'percent','美国财政部原生10年平价实际收益率，非通胀预期或名义收益率')
 spread=[{'date':d,'value':cn[d]['value']-nominal[d]['value'],'cn_10y':cn[d]['value'],'us_10y':nominal[d]['value'],'source_ids':cn[d]['source_ids']+nominal[d]['source_ids']} for d in sorted(cn.keys()&nominal.keys()) if d>='2025-01-01']
 if spread:add('中美10年国债利差',spread,'percentage_points','中国财政部中国国债收益率曲线10年 − 美国财政部平价国债收益率10年；仅两源同日期交集，不填补休市差异')
 if 'bea_native_nipa' in names:
  wb=load_workbook(io.BytesIO(read('bea_native_nipa')),read_only=True,data_only=True);sh=wb['T20804-M'];rows=list(sh.iter_rows(values_only=True));header=rows[7];matches=[r for r in rows if len(r)>2 and r[2]=='DPCCRG']
  if len(matches)!=1 or '2017=100' not in str(rows[1]):raise ValueError('wrong_native_core_pce_identity')
  clock=re.search(r'Data published (\w+ \d{1,2}, \d{4})',str(rows[4]));release=datetime.strptime(clock[1],'%B %d, %Y').date().isoformat();points=[]
  for i,label in enumerate(header):
   if isinstance(label,str) and re.fullmatch(r'202[5-9]M\d{2}',label):
    y,m=int(label[:4]),int(label[5:]);d=f'{y}-{m:02}-{calendar.monthrange(y,m)[1]}'
    if d<=as_of_date:points.append({'date':d,'value':float(matches[0][i]),'period':label,'release_date':release,'source_ids':['bea_native_nipa']})
  add('核心PCE价格指数',points,'BEA_core_PCE_2017_100','BEA Table 2.8.4月度 DPCCRG，剔除食品能源；季调、2017=100；当前完整修订版本公布于'+release+'，不是各月历史首发值')
 economic={'工业增加值同比':[],'社零同比':[],'固定资产投资累计同比':[],'GDP同比':[],'M2同比':[]}
 for s in sources:
  sid=s['id']
  if not(sid.startswith('nbs_gdp_') or sid.startswith('nbs_econ_') or sid.startswith('pbc_m2_') or sid=='pbc_jan_official_reprint'):continue
  p,t=text(read(sid));release=p.publication_date()
  if release is None:
   clock=re.search(r'202[5-9][/\-]\d{2}[/\-]\d{2}',t);release=clock[0].replace('/','-') if clock else None
  # Original report year/month, not URL publication month.
  patterns=[('工业增加值同比',r'(202[5-9])年(?:(1[—－-]2)|([1-9]|1[0-2]))月份(?:全国)?规模以上工业增加值(?:同比)?(?:实际)?(?:增长|下降)([\d.]+)%'),('社零同比',r'(202[5-9])年(?:(1[—－-]2)|([1-9]|1[0-2]))月份社会消费品零售总额(?:同比)?(?:增长|下降)([\d.]+)%'),('固定资产投资累计同比',r'(202[5-9])年1[—－-](\d{1,2})月份全国固定资产投资(?:[（(]不含农户[）)])?(?:同比)?(?:增长|下降)([\d.]+)%')]
  for metric,pat in patterns:
   m=re.search(pat,t)
   if m:
    y=int(m[1]);month=int(m[2]) if metric=='固定资产投资累计同比' else 2 if m[2] else int(m[3]);value=float(m.group(m.lastindex));sign=-1 if '下降' in m[0] else 1;d=f'{y}-{month:02}-{calendar.monthrange(y,month)[1]}'
    if release and release<=as_of_date:economic[metric].append({'date':d,'value':value*sign,'release_date':release,'source_ids':[sid],'period_basis':'January-February combined' if month==2 else 'YTD' if metric=='固定资产投资累计同比' else 'monthly'})
  # Basic-situation titles omit growth; read the native total from the first body paragraph.
  metric=next((name for name,needle in [('固定资产投资累计同比','全国固定资产投资'),('社零同比','社会消费品零售总额'),('工业增加值同比','规模以上工业增加值')] if needle in t[:130]),None)
  if metric and not any(sid in r['source_ids'] for r in economic[metric]):
   identity=re.search(r'(202[5-9])年(?:(?:1[—－-])?(\d{1,2})月份|上半年|全年)?',t[:160])
   if metric=='固定资产投资累计同比':pat=r'全国固定资产投资[（(]不含农户[）)]\s*[\d.]+亿元，(?:同比|比上年)(增长|下降)([\d.]+)%'
   elif metric=='社零同比':pat=r'社会消费品零售总额\s*[\d.]+亿元，(?:同比|比上年)(增长|下降)([\d.]+)%'
   else:pat=r'规模以上工业增加值同比(?:实际)?(增长|下降)([\d.]+)%'
   native=re.search(pat,t[:1600])
   if identity and native and release and release<=as_of_date:
    y=int(identity[1]);month=int(identity[2]) if identity[2] else 6 if '上半年' in identity[0] else 12;value=float(native[2])*(-1 if native[1]=='下降' else 1);d=f'{y}-{month:02}-{calendar.monthrange(y,month)[1]}'
    if d<=as_of_date:economic[metric].append({'date':d,'value':value,'release_date':release,'source_ids':[sid],'period_basis':'January-February combined' if month==2 else 'YTD' if metric=='固定资产投资累计同比' else 'monthly'})
  # Each quarter uses its own initial native Table 1, avoiding later revised history.
  if sid.startswith('nbs_gdp_') and release and release<=as_of_date:
   identity=re.search(r'(202[5-9])年(一|二|三|四)季度',t);gdp=re.search(r'GDP ([\d.]+) ([\d.]+)(?: ([\d.]+) ([\d.]+))? 第一产业',t)
   if identity and gdp:
    y=int(identity[1]);month={'一':3,'二':6,'三':9,'四':12}[identity[2]];value=float(gdp[3] if gdp[3] else gdp[2]);d=f'{y}-{month:02}-{calendar.monthrange(y,month)[1]}'
    if d<=as_of_date:economic['GDP同比'].append({'date':d,'value':value,'release_date':release,'source_ids':[sid],'period_basis':'native quarter YoY initial Table 1'})
  if sid.startswith('pbc_'):
   m=re.search(r'广义货币[（(]M2[）)]余额[^。]{0,100}?同比增长\s*([\d.]+)',t)
   identity=re.search(r'pbc_m2_yoy_(202[5-9])_(\d{2})',sid)
   if sid=='pbc_jan_official_reprint':identity=re.match(r'(2025)_(01)','2025_01');release='2025-02-17'
   if m and identity and release:
    y,month=int(identity[1]),int(identity[2]);economic['M2同比'].append({'date':f'{y}-{month:02}-{calendar.monthrange(y,month)[1]}','value':float(m[1]),'release_date':release,'source_ids':[sid],'fallback_status':'Shanghai government repost of PBC original; original 404' if sid=='pbc_jan_official_reprint' else 'none'})
 for name,points in economic.items():
  if points:add(name,points,'percent',{'GDP同比':'GDP原生当季同比：各季初步核算Table 1当季列；保留各季原始发布版本，不拼入后来修订表' ,'固定资产投资累计同比':'固定资产投资（不含农户）累计同比，1—2月合并，其后累计期间','工业增加值同比':'规模以上工业增加值当月同比，1—2月合并','社零同比':'社会消费品零售总额当月同比，1—2月合并','M2同比':'中国广义货币月末余额同比；2025年1月原始页404，政府转载显式标注'}[name])
 try:
  facts={b:json.loads(read('sec_'+b)) for b in ('amzn','msft','goog')}
  if any(facts[b].get('cik')!=cik for b,cik in {'amzn':1018724,'msft':789019,'goog':1652044}.items()):raise ValueError('wrong_native_sec_cik')
  dc=dc_rows(sources,read);ai=compute(facts,dc,real,as_of_date)
  add('AI压力',[{'date':r['date'],'value':r['total'],'release_date':r['release_date'],'source_ids':r['source_ids'],'contributions':r['contributions'],'scores':r['scores']} for r in ai],'pressure_0_100','自建公开财务压力；需求40%=NVIDIA数据中心TTM同比（反向）；现金流30%=三大买家Capex/OCF；外融25%=原生现金债务发行/Capex；约束5%=10年TIPS。按披露日阶梯更新；前12个不同季度最少8个及前252个收益率最少60个，用median/1.4826MAD标准化；缺项不重新分配权重。',step=True)
 except (KeyError,ValueError) as e:ai=[];dc=[];errors['AI压力']=str(e)
 existing={s['name'] for s in series if any(r['value'] is not None for r in s['points'])};gaps=[{'chart':cid,'metric':name,'reason':errors.get(name,'原生序列尚未取得或四权重有效历史不足')} for cid,_,metrics in CHARTS for name in metrics if name not in existing]
 packet={'schema_version':'five-charts-v1','scope_key':scope_key,'run_id':run_id,'as_of_date':as_of_date,'generated_at':generated_at or datetime.now(timezone.utc).isoformat(),'action':'NO_ACTION','historical_vintage_certified':False,'sources':sources,'series':series,'gaps':gaps,'pressure_events':ai,'dc_quarters':dc,'notes':notes}
 return packet
